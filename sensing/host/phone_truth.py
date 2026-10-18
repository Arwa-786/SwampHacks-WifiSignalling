"""True breathing rate from a phone on the belly (phyphox), to grade the Wi-Fi.

The same idea as Gu et al. 2019 (an accelerometer on the abdomen as ground
truth), with a free app instead of a sensor: phyphox, "Acceleration (without g)",
phone flat against the belly, screen out. Breathing pushes the belly forward and
back, which the accelerometer sees as a slow wave at the breathing rate.

Export from phyphox as Excel (.xlsx) or CSV (.zip); both carry the clock time the
recording started ("Metadata Time"), and the Wi-Fi recorders write the offset
between the laptop's uptime clock and real time into their header (clock=...),
so the two line up to well under a second.
"""

from __future__ import annotations

import csv
import io
import os
import re
import zipfile

import numpy as np
from scipy.signal import butter, sosfiltfilt

# Hz = 8-42 breaths per minute. Starting at 0.1 Hz let slow drift win: our first phone
# test read 6/min on one axis while the other two agreed on 19/min.
BAND = (0.13, 0.7)


def load_phyphox(path: str) -> tuple[np.ndarray, np.ndarray]:
    """(clock time of each sample in seconds since 1970, acceleration n x 3)."""
    if path.endswith(".xlsx"):
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True)
        raw = list(wb["Raw Data"].iter_rows(values_only=True))[1:]
        events = list(wb["Metadata Time"].iter_rows(values_only=True))[1:]
    elif path.endswith(".zip"):
        with zipfile.ZipFile(path) as z:
            def table(name):
                member = next(n for n in z.namelist() if n.endswith(name))
                return list(csv.reader(io.TextIOWrapper(z.open(member), "utf-8")))[1:]
            raw, events = table("Raw Data.csv"), table("time.csv")
    else:
        raise SystemExit(f"{path}: export from phyphox as Excel (.xlsx) or CSV (.zip)")
    data = np.array([[float(v) for v in r[:4]] for r in raw if r and r[0] not in (None, "")])
    # each START/PAUSE event pairs the experiment's own timer with the clock time
    starts = [(float(e[1]), float(e[2])) for e in events if e[0] == "START"]
    exp_t = data[:, 0]
    offset = np.zeros_like(exp_t)
    for k, (et, wall) in enumerate(starts):  # a later START (after a pause) re-anchors what follows
        later = np.ones(len(exp_t), bool) if k == 0 else exp_t >= et - 1e-6
        offset[later] = wall - et
    return exp_t + offset, data[:, 1:4]


def clock_offset(log_path: str) -> float:
    """Real time minus laptop uptime for a Wi-Fi recording (from its header)."""
    with open(log_path) as f:
        header = f.readline()
    m = re.search(r"clock=([0-9.]+)", header)
    if not m:
        raise SystemExit(f"{os.path.basename(log_path)} has no clock= in its header: record it again "
                         "with the current recorder to line it up with the phone")
    return float(m.group(1))


def rate(times: np.ndarray, acc: np.ndarray, t0: float, t1: float) -> tuple[float, float]:
    """(breaths per minute, rhythm clarity 0-1) over [t0, t1], from whichever
    axis shows the clearest rhythm (how the phone sits on the belly varies)."""
    m = (times >= t0) & (times < t1)
    if m.sum() < 20 * (t1 - t0):  # fewer than ~20 samples/s: gap in the phone recording
        return float("nan"), 0.0
    t, a = times[m], acc[m]
    fs = 1 / np.median(np.diff(t))
    best = (float("nan"), 0.0)
    for axis in range(3):
        x = a[:, axis] - a[:, axis].mean()
        x = sosfiltfilt(butter(2, BAND, btype="bandpass", fs=fs, output="sos"), x)
        n = 16 * len(x)
        spec = np.abs(np.fft.rfft(x * np.hanning(len(x)), n)) ** 2
        f = np.fft.rfftfreq(n, 1 / fs)
        band = (f >= BAND[0]) & (f <= BAND[1])
        peak = f[band][np.argmax(spec[band])]
        clarity = float(spec[band][np.abs(f[band] - peak) <= 0.05].sum() / (spec[band].sum() + 1e-12))
        if clarity > best[1]:
            best = (float(peak * 60), clarity)
    return best
